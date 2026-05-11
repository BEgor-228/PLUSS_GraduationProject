import json
import re
import csv
from datetime import datetime, timedelta

from django.contrib.auth import authenticate, get_user_model, login as auth_login, logout as auth_logout
from django.contrib.auth.models import Group
from django.db import IntegrityError, connection, transaction
from django.db.models import Avg, Count, Q
from django.db.models.functions import TruncMonth
from django.http import HttpRequest, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_GET, require_POST

from .models import (
    ActionLog,
    AccessibilityCriterionType,
    AdmissionType,
    AoopProgram,
    ConditionType,
    Director,
    District,
    Favorite,
    Institution,
    InstitutionReview,
    InstitutionType,
    Link,
    ProfileNotification,
    UserBlockState,
    UserDistrictPreference,
)


BAYESIAN_CONFIDENCE_THRESHOLD = 10


def index(request: HttpRequest):
    return render(request, "mapapp/index.html")


def profile_page(request: HttpRequest):
    user = request.user
    ctx: dict = {"profile_role": "guest"}
    if user.is_authenticated:
        is_admin = user.is_superuser or user.groups.filter(name="administrators").exists()
        login_base = (user.username or "").split("@")[0].strip()
        initials_source = login_base or (user.get_full_name() or user.email or "U")
        letters = [ch for ch in initials_source if ch.isalnum()]
        profile_initials = "".join(letters[:2]).upper() if letters else "U"
        ctx = {
            "profile_role": "admin" if is_admin else "portal",
            "display_name": user.get_full_name() or user.username,
            "display_email": user.email,
            "profile_initials": profile_initials,
        }
        if is_admin:
            now = timezone.now()
            month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            total_institutions = Institution.objects.count()
            institutions_month_delta = Institution.objects.filter(created_at__gte=month_start).count()
            User = get_user_model()
            total_users = User.objects.count()
            users_month_delta = User.objects.filter(date_joined__gte=month_start).count()
            moderation_total = InstitutionReview.objects.filter(
                status__in=[
                    InstitutionReview.STATUS_PENDING,
                    InstitutionReview.STATUS_DISAPPROVED,
                    "rejected",
                ]
            ).count()
            moderation_month_delta = InstitutionReview.objects.filter(
                status__in=[
                    InstitutionReview.STATUS_PENDING,
                    InstitutionReview.STATUS_DISAPPROVED,
                    "rejected",
                ],
                created_at__gte=month_start,
            ).count()
            type_distribution_qs = (
                Institution.objects.values("type__name_ru")
                .annotate(total=Count("id"))
                .order_by("-total")
            )
            type_distribution = [
                {
                    "name": row["type__name_ru"] or "Без типа",
                    "count": row["total"],
                }
                for row in type_distribution_qs
            ]
            district_distribution_qs = (
                Institution.objects.values("district__name")
                .annotate(total=Count("id"))
                .order_by("-total")
            )
            district_distribution = [
                {
                    "name": row["district__name"] or "Район не указан",
                    "count": row["total"],
                }
                for row in district_distribution_qs
            ]
            user_reviews_chart_qs = (
                User.objects.filter(is_superuser=False)
                .annotate(reviews_count=Count("institution_reviews"))
                .filter(reviews_count__gt=0)
                .order_by("-reviews_count", "username")
                .values("first_name", "last_name", "username", "reviews_count")
            )
            user_reviews_chart = []
            for row in user_reviews_chart_qs:
                full_name = " ".join(part for part in [row["first_name"], row["last_name"]] if part).strip()
                display_name = full_name or row["username"] or "Пользователь"
                user_reviews_chart.append(
                    {
                        "name": display_name,
                        "username": row["username"] or "",
                        "value": int(row["reviews_count"] or 0),
                    }
                )
            global_rating_mean = _get_global_approved_rating_mean()
            institution_rating_chart_qs = (
                Institution.objects.annotate(
                    avg_rating=Avg(
                        "reviews__rating",
                        filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED),
                    ),
                    approved_reviews_count=Count(
                        "reviews",
                        filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED),
                    ),
                )
                .select_related("district")
                .order_by("name")
            )
            institution_rating_chart = []
            for inst in institution_rating_chart_qs:
                bayesian_value = _calc_bayesian_rating(
                    float(inst.avg_rating) if inst.avg_rating is not None else None,
                    int(inst.approved_reviews_count or 0),
                    global_rating_mean,
                )
                # Учреждения без одобренных отзывов тоже показываем в рейтинге дашборда.
                score_value = bayesian_value if bayesian_value is not None else 0.0
                institution_rating_chart.append(
                    {
                        "name": inst.name,
                        "district_name": inst.district.name if inst.district_id else "",
                        "value": score_value,
                    }
                )
            institution_rating_chart.sort(key=lambda row: (-float(row.get("value") or 0), row.get("name") or ""))
            activity_qs = (
                User.objects.filter(is_superuser=False)
                .annotate(month=TruncMonth("date_joined"))
                .values("month")
                .annotate(c=Count("id"))
                .order_by("month")
            )
            admin_user_activity_monthly = [
                {"month": row["month"].strftime("%Y-%m"), "count": int(row["c"] or 0)}
                for row in activity_qs
                if row.get("month")
            ]
            admin_users_qs = (
                User.objects.filter(is_superuser=False)
                .annotate(
                    reviews_count=Count("institution_reviews"),
                    avg_rating=Avg("institution_reviews__rating"),
                )
                .order_by("-date_joined")
                .values(
                    "id",
                    "username",
                    "first_name",
                    "last_name",
                    "email",
                    "reviews_count",
                    "avg_rating",
                    "date_joined",
                )
            )
            admin_user_rows = list(admin_users_qs)
            block_by_uid = {
                b.user_id: b
                for b in UserBlockState.objects.filter(
                    user_id__in=[r["id"] for r in admin_user_rows],
                )
            }
            admin_users = []
            for row in admin_user_rows:
                full_name = " ".join(part for part in [row["first_name"], row["last_name"]] if part).strip()
                display_name = full_name or row["username"] or "Пользователь"
                initials_letters = [ch for ch in display_name if ch.isalnum()]
                initials = "".join(initials_letters[:2]).upper() if initials_letters else "U"
                bs = block_by_uid.get(row["id"])
                is_blocked = bool(bs and bs.is_blocked)
                admin_users.append(
                    {
                        "id": row["id"],
                        "display_name": display_name,
                        "username": row["username"] or "",
                        "email": row["email"] or "",
                        "initials": initials,
                        "is_current_user": row["id"] == user.id,
                        "reviews_count": int(row.get("reviews_count") or 0),
                        "avg_rating": float(row["avg_rating"]) if row.get("avg_rating") is not None else None,
                        "date_joined": row["date_joined"],
                        "is_blocked": is_blocked,
                        "block_reason": (bs.block_reason or "").strip() if bs else "",
                    }
                )
            admin_institutions_qs = (
                Institution.objects.annotate(
                    avg_rating=Avg(
                        "reviews__rating",
                        filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED),
                    ),
                    approved_reviews_count=Count(
                        "reviews",
                        filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED),
                    ),
                    total_reviews_count=Count("reviews"),
                )
                .select_related("district", "type")
                .order_by("name")
            )
            admin_institutions = []
            for inst in admin_institutions_qs:
                display_name = inst.name or "Учреждение"
                initials_source = display_name.strip() or "У"
                letters_inst = [ch for ch in initials_source if ch.isalnum()]
                inst_initials = "".join(letters_inst[:2]).upper() if letters_inst else "У"
                admin_institutions.append(
                    {
                        "id": inst.id,
                        "name": display_name,
                        "district_name": inst.district.name if inst.district_id else "Район не указан",
                        "type_name": inst.type.name_ru if inst.type_id else "Без типа",
                        "initials": inst_initials,
                        "avg_rating": _calc_bayesian_rating(
                            float(inst.avg_rating) if inst.avg_rating is not None else None,
                            int(inst.approved_reviews_count or 0),
                            global_rating_mean,
                        ),
                        "approved_reviews_count": int(inst.approved_reviews_count or 0),
                        "total_reviews_count": int(inst.total_reviews_count or 0),
                    }
                )
            moderation_reviews_qs = (
                InstitutionReview.objects.filter(
                    status__in=[
                        InstitutionReview.STATUS_PENDING,
                        InstitutionReview.STATUS_DISAPPROVED,
                        "rejected",
                    ]
                )
                .select_related("institution__type", "user")
                .order_by("-created_at")
            )
            moderation_reviews = [
                {
                    "id": review.id,
                    "institution_type": review.institution.type.name_ru if review.institution and review.institution.type else "Без типа",
                    "institution_name": review.institution.name if review.institution else "Учреждение",
                    "user_name": review.user.get_full_name() or review.user.username or "Пользователь",
                    "status_label": (
                        "На модерации"
                        if review.status == InstitutionReview.STATUS_PENDING
                        else "Отклонен"
                    ),
                    "status_value": (
                        review.status
                        if review.status in {InstitutionReview.STATUS_PENDING, InstitutionReview.STATUS_DISAPPROVED}
                        else InstitutionReview.STATUS_DISAPPROVED
                    ),
                    "rating": float(review.rating or 0),
                    "comment": review.comment or "Комментарий не указан.",
                    "created_at": review.created_at,
                }
                for review in moderation_reviews_qs
            ]
            ctx.update(
                {
                    "admin_total_institutions": total_institutions,
                    "admin_institutions_month_delta": institutions_month_delta,
                    "admin_total_users": total_users,
                    "admin_users_month_delta": users_month_delta,
                    "admin_moderation_total": moderation_total,
                    "admin_moderation_month_delta": moderation_month_delta,
                    "admin_type_distribution": type_distribution,
                    "admin_district_distribution": district_distribution,
                    "admin_user_reviews_chart": user_reviews_chart,
                    "admin_user_activity_monthly": admin_user_activity_monthly,
                    "admin_institution_rating_chart": institution_rating_chart,
                    "admin_users": admin_users,
                    "admin_institutions": admin_institutions,
                    "admin_moderation_reviews": moderation_reviews,
                }
            )
        if not is_admin:
            ctx["settings_first_name"] = user.first_name or ""
            ctx["settings_username"] = user.username or ""
            ctx["settings_email"] = user.email or ""
            preference = (
                UserDistrictPreference.objects.filter(user=user)
                .select_related("district", "preferred_institution_type")
                .first()
            )
            pref = preference
            ctx["settings_notify_new_institutions"] = bool(pref and pref.notify_new_institutions)
            ctx["settings_districts"] = list(District.objects.order_by("name").values("id", "name"))
            ctx["main_pref_district_id"] = pref.district_id if pref else None
            ctx["main_pref_institution_type"] = pref.preferred_institution_type_id if pref else None
            ctx["main_pref_child_age_min"] = pref.child_age_min if pref else None
            ctx["main_pref_child_age_max"] = pref.child_age_max if pref else None
            ctx["main_pref_school_class_from"] = pref.school_class_from if pref else None
            ctx["main_pref_school_class_to"] = pref.school_class_to if pref else None
            ctx["main_pref_condition_codes"] = list(pref.preferred_condition_codes or []) if pref else []
            ctx["main_pref_accessibility_codes"] = list(pref.preferred_accessibility_codes or []) if pref else []
            ctx["main_pref_prefer_aoop"] = bool(pref and pref.prefer_aoop)
            ctx["main_institution_types"] = list(InstitutionType.objects.order_by("name_ru").values("code", "name_ru"))
            ctx["main_condition_types"] = list(ConditionType.objects.order_by("name_ru").values("code", "name_ru"))
            ctx["main_accessibility_types"] = list(
                AccessibilityCriterionType.objects.order_by("name_ru").values("code", "name_ru")
            )
            favorites = (
                Favorite.objects.filter(user=user)
                .select_related("institution__district", "institution__type")
                .order_by("-created_at")
            )
            favorite_cards = []
            for fav in favorites:
                inst = fav.institution
                favorite_cards.append(
                    {
                        "id": inst.id,
                        "district_id": inst.district_id,
                        "name": inst.name,
                        "district_name": inst.district.name if inst.district_id else "",
                        "type_name": inst.type.name_ru if inst.type_id else "",
                        "range_min": inst.range_min,
                        "range_max": inst.range_max,
                    }
                )
            ctx["favorite_institutions"] = favorite_cards
            ctx["favorites_count"] = len(favorite_cards)
            user_reviews_qs = (
                InstitutionReview.objects.filter(user=user)
                .select_related("institution__district", "institution__type")
                .order_by("-created_at")
            )
            status_label_map = {
                InstitutionReview.STATUS_PENDING: "На модерации",
                InstitutionReview.STATUS_APPROVED: "Опубликован",
                InstitutionReview.STATUS_DISAPPROVED: "Отклонен",
                "rejected": "Отклонен",
            }
            user_reviews = []
            for review in user_reviews_qs:
                user_reviews.append(
                    {
                        "id": review.id,
                        "institution_name": review.institution.name if review.institution_id else "Учреждение",
                        "district_name": (
                            review.institution.district.name
                            if review.institution_id and review.institution.district_id
                            else "Город не указан"
                        ),
                        "type_name": (
                            review.institution.type.name_ru
                            if review.institution_id and review.institution.type_id
                            else "Тип не указан"
                        ),
                        "comment": review.comment or "Комментарий не указан.",
                        "rating": float(review.rating or 0),
                        "status_value": (
                            review.status
                            if review.status in {
                                InstitutionReview.STATUS_PENDING,
                                InstitutionReview.STATUS_APPROVED,
                                InstitutionReview.STATUS_DISAPPROVED,
                            }
                            else InstitutionReview.STATUS_DISAPPROVED
                        ),
                        "status_label": status_label_map.get(review.status, "На модерации"),
                        "created_at": review.created_at,
                        "created_at_display": review.created_at.strftime("%d.%m.%Y %H:%M"),
                    }
                )
            ctx["user_reviews"] = user_reviews
            ctx["user_reviews_count"] = len(user_reviews)
            notifications_qs = (
                ProfileNotification.objects.filter(user=user)
                .select_related("institution__district", "institution__type")
                .order_by("-created_at")
            )
            notifications = []
            for notification in notifications_qs:
                inst = notification.institution
                notifications.append(
                    {
                        "institution_id": inst.id if inst else None,
                        "district_id": inst.district_id if inst else None,
                        "district_name": (
                            inst.district.name if inst and inst.district_id else "Город/район не указан"
                        ),
                        "type_name": (
                            inst.type.name_ru if inst and inst.type_id else "Тип не указан"
                        ),
                        "institution_name": inst.name if inst else "Учреждение",
                        "message": notification.message,
                        "created_at": notification.created_at,
                    }
                )
            ctx["notifications"] = notifications
            ctx["notifications_count"] = len(notifications)
    return render(request, "mapapp/profile.html", ctx)


def district_page(request: HttpRequest, district_id: int):
    district = get_object_or_404(District, pk=district_id)
    return render(request, "mapapp/district.html", {"district": district})


@require_GET
def institution_create_page(request: HttpRequest):
    if not _require_admin(request):
        return redirect("index")

    district_id = request.GET.get("district_id")
    initial_data = {"district_id": int(district_id)} if district_id and district_id.isdigit() else {}
    districts = District.objects.order_by("name").values("id", "name")
    return render(
        request,
        "mapapp/institution_form.html",
        {"initial_data": initial_data, "districts": districts, "mode": "create"},
    )


@require_GET
def institution_edit_page(request: HttpRequest, institution_id: int):
    if not _require_admin(request):
        return redirect("index")

    inst = get_object_or_404(
        Institution.objects.select_related("director", "type").prefetch_related(
            "conditions", "admission", "aoop_programs"
        ),
        pk=institution_id,
    )
    districts = District.objects.order_by("name").values("id", "name")
    return render(
        request,
        "mapapp/institution_form.html",
        {"initial_data": _serialize_institution(inst), "districts": districts, "mode": "edit"},
    )


def _parse_json(request: HttpRequest):
    try:
        return json.loads(request.body.decode("utf-8")) if request.body else {}
    except json.JSONDecodeError:
        return None


def _get_pagination(request: HttpRequest) -> tuple[int, int]:
    try:
        page = int(request.GET.get("page", "1"))
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = int(request.GET.get("page_size", "10"))
    except (TypeError, ValueError):
        page_size = 10
    page = max(1, page)
    page_size = max(1, min(page_size, 100))
    return page, page_size


def _build_pagination(total_items: int, page: int, page_size: int) -> dict:
    total_pages = max(1, (total_items + page_size - 1) // page_size)
    safe_page = min(max(1, page), total_pages)
    return {
        "page": safe_page,
        "page_size": page_size,
        "total_items": total_items,
        "total_pages": total_pages,
    }


def _user_block_row(user):
    if not user or not user.is_authenticated:
        return None
    return UserBlockState.objects.filter(user_id=user.pk).first()


def _user_is_blocked(user) -> bool:
    row = _user_block_row(user)
    return bool(row and row.is_blocked)


def _user_block_reason(user) -> str:
    row = _user_block_row(user)
    if not row or not row.is_blocked:
        return ""
    return (row.block_reason or "").strip()


def _blocked_login_response(user):
    return JsonResponse(
        {
            "success": False,
            "blocked": True,
            "reason": _user_block_reason(user),
            "error": "Аккаунт заблокирован",
        },
        status=403,
    )


def _authenticate_with_login_or_email(request: HttpRequest, login_value: str, password: str):
    User = get_user_model()
    login_value = (login_value or "").strip()
    if not login_value:
        return None
    user = authenticate(request, username=login_value, password=password)
    if user:
        return user
    by_email = User.objects.filter(email__iexact=login_value).first()
    if by_email:
        return authenticate(request, username=by_email.username, password=password)
    return None


def _repair_auth_user_pk_sequence():
    if connection.vendor != "postgresql":
        return
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT setval(
                pg_get_serial_sequence('auth_user', 'id'),
                COALESCE((SELECT MAX(id) FROM auth_user), 1),
                (SELECT MAX(id) IS NOT NULL FROM auth_user)
            )
            """
        )


def _create_user_with_repaired_sequence(**kwargs):
    User = get_user_model()
    try:
        return User.objects.create_user(**kwargs)
    except IntegrityError as exc:
        if "auth_user_pkey" not in str(exc):
            raise
        _repair_auth_user_pk_sequence()
        return User.objects.create_user(**kwargs)


def _get_global_approved_rating_mean() -> float:
    value = (
        InstitutionReview.objects.filter(status=InstitutionReview.STATUS_APPROVED)
        .aggregate(v=Avg("rating"))
        .get("v")
    )
    return float(value) if value is not None else 0.0


def _calc_bayesian_rating(avg_rating: float | None, rating_count: int, global_mean: float) -> float | None:
    if avg_rating is None or rating_count <= 0:
        return None
    v = float(rating_count)
    m = float(BAYESIAN_CONFIDENCE_THRESHOLD)
    weighted = ((v / (v + m)) * float(avg_rating)) + ((m / (v + m)) * global_mean)
    return round(weighted, 2)


def _get_reviews_aggregates_for_institutions(institution_ids: list[int]) -> dict[int, dict]:
    if not institution_ids:
        return {}

    reviews = InstitutionReview.objects.filter(
        institution_id__in=institution_ids,
        status=InstitutionReview.STATUS_APPROVED,
    ).only(
        "institution_id", "rating", "criteria_scores"
    )

    inst_rating_sum: dict[int, float] = {}
    inst_rating_count: dict[int, int] = {}

    criteria_sum: dict[int, dict[str, float]] = {}
    criteria_count: dict[int, dict[str, int]] = {}

    for review in reviews:
        inst_id = review.institution_id
        inst_rating_sum[inst_id] = inst_rating_sum.get(inst_id, 0.0) + float(review.rating or 0)
        inst_rating_count[inst_id] = inst_rating_count.get(inst_id, 0) + 1

        criteria_scores = review.criteria_scores or {}
        if isinstance(criteria_scores, dict):
            for code, score in criteria_scores.items():
                try:
                    score_int = int(score)
                except (TypeError, ValueError):
                    continue
                code_s = str(code)
                criteria_sum.setdefault(inst_id, {}).setdefault(code_s, 0.0)
                criteria_count.setdefault(inst_id, {}).setdefault(code_s, 0)
                criteria_sum[inst_id][code_s] += float(score_int)
                criteria_count[inst_id][code_s] += 1

    global_mean = _get_global_approved_rating_mean()
    aggregates: dict[int, dict] = {}
    for inst_id in institution_ids:
        rating_avg_raw = None
        rating_count = int(inst_rating_count.get(inst_id) or 0)
        if rating_count:
            rating_avg_raw = inst_rating_sum[inst_id] / rating_count
        rating_avg = _calc_bayesian_rating(rating_avg_raw, rating_count, global_mean)

        avg_criteria: dict[str, float] = {}
        codes_for_inst = criteria_sum.get(inst_id, {})
        for code, sum_value in codes_for_inst.items():
            cnt = criteria_count.get(inst_id, {}).get(code, 0)
            if cnt:
                avg_criteria[code] = round(sum_value / cnt, 2)

        aggregates[inst_id] = {
            "avg_institution_rating": rating_avg,
            "avg_accessibility_criteria": avg_criteria,
        }

    return aggregates


def _get_approved_reviews_for_institutions(
    institution_ids: list[int], limit_per_institution: int = 3
) -> dict[int, dict]:
    if not institution_ids:
        return {}
    counts_qs = (
        InstitutionReview.objects.filter(
            institution_id__in=institution_ids,
            status=InstitutionReview.STATUS_APPROVED,
        )
        .values("institution_id")
        .annotate(total=Count("id"))
    )
    counts_map = {row["institution_id"]: row["total"] for row in counts_qs}
    result: dict[int, dict] = {
        inst_id: {"count": counts_map.get(inst_id, 0), "items": []} for inst_id in institution_ids
    }
    reviews_qs = (
        InstitutionReview.objects.filter(
            institution_id__in=institution_ids,
            status=InstitutionReview.STATUS_APPROVED,
        )
        .select_related("user")
        .order_by("institution_id", "-created_at", "-id")
    )
    per_inst_added: dict[int, int] = {}
    for review in reviews_qs:
        inst_id = review.institution_id
        current = per_inst_added.get(inst_id, 0)
        if current >= limit_per_institution:
            continue
        user_name = review.user.get_full_name() or review.user.username or "Пользователь"
        result.setdefault(inst_id, {"count": 0, "items": []})
        result[inst_id]["items"].append(
            {
                "author": user_name,
                "rating": float(review.rating or 0),
                "comment": review.comment or "",
                "created_at": review.created_at.strftime("%d.%m.%Y %H:%M"),
            }
        )
        per_inst_added[inst_id] = current + 1
    return result


def _serialize_institution(
    inst: Institution,
    favorite_ids: set[int] | None = None,
    review_aggs: dict[int, dict] | None = None,
    approved_reviews_map: dict[int, dict] | None = None,
):
    is_favorite = bool(favorite_ids and inst.id in favorite_ids)
    review_agg = (review_aggs or {}).get(inst.id) if review_aggs else None
    approved_reviews_payload = (approved_reviews_map or {}).get(inst.id, {})
    conditions_ordered = sorted(inst.conditions.all(), key=lambda c: (c.name_ru or "").lower())
    admission_ordered = sorted(inst.admission.all(), key=lambda a: (a.name_ru or "").lower())
    accessibility_ordered = sorted(
        inst.accessibility_criteria.all(), key=lambda a: (a.name_ru or "").lower()
    )
    return {
        "id": inst.id,
        "name": inst.name,
        "address": inst.address,
        "description": inst.description,
        "range_min": inst.range_min,
        "range_max": inst.range_max,
        "website": inst.website,
        "aoop_url": inst.aoop_url,
        "district_id": inst.district_id,
        "district_name": inst.district.name if inst.district_id else "",
        "type": inst.type_id,
        "type_name_ru": inst.type.name_ru if inst.type_id else "",
        "conditions": list(inst.conditions.values_list("code", flat=True)),
        "condition_names_ru": [c.name_ru for c in conditions_ordered],
        "conditionsAdmission": list(inst.admission.values_list("code", flat=True)),
        "admission_names_ru": [a.name_ru for a in admission_ordered],
        "accessibility_criteria": list(inst.accessibility_criteria.values_list("code", flat=True)),
        "accessibility_names_ru": [a.name_ru for a in accessibility_ordered],
        "aoop_programs": [{"id": p.id, "name": p.name} for p in inst.aoop_programs.all()],
        "is_favorite": is_favorite,
        "avg_institution_rating": review_agg.get("avg_institution_rating") if review_agg else None,
        "avg_accessibility_criteria": (review_agg.get("avg_accessibility_criteria") if review_agg else {}) or {},
        "approved_reviews_count": approved_reviews_payload.get("count", 0),
        "approved_reviews": approved_reviews_payload.get("items", []),
        "director": {
            "id": inst.director.id if inst.director else None,
            "name": inst.director.full_name if inst.director else "",
            "phone": inst.director.phone if inst.director else "",
            "email": inst.director.email if inst.director else "",
        },
    }


def _get_favorite_ids_for_request(request: HttpRequest) -> set[int]:
    user = request.user
    if not user.is_authenticated:
        return set()
    if user.is_superuser or user.groups.filter(name="administrators").exists():
        return set()
    return set(Favorite.objects.filter(user=user).values_list("institution_id", flat=True))


def _normalize_search_value(value: str) -> str:
    if not value:
        return ""
    value = value.lower().replace("ё", "е").replace("№", "")
    return re.sub(r"[^a-zа-я0-9]", "", value)


def _matches_age_bucket(inst_payload: dict, age_bucket: str) -> bool:
    range_min = inst_payload.get("range_min")
    range_max = inst_payload.get("range_max")
    if range_min is None or range_max is None:
        return False
    if "-" in age_bucket:
        min_str, max_str = age_bucket.split("-", 1)
        try:
            min_value = float(min_str)
            max_value = float(max_str)
        except ValueError:
            return False
        return range_min <= max_value and range_max >= min_value
    if "+" in age_bucket:
        try:
            min_value = int(age_bucket.replace("+", ""))
        except ValueError:
            return False
        return range_max >= min_value
    return False


def _age_year_to_buckets(age_int: int) -> set[str]:
    buckets: set[str] = set()
    if age_int <= 3:
        buckets.add("1.5-3")
    if 3 <= age_int <= 5:
        buckets.add("3-5")
    if 5 <= age_int <= 7:
        buckets.add("5-7")
    if age_int >= 7:
        buckets.add("7+")
    return buckets


def _years_span_to_age_buckets(age_min: int | None, age_max: int | None) -> list[str]:
    if age_min is None and age_max is None:
        return []
    lo = int(age_min) if age_min is not None else 0
    hi = int(age_max) if age_max is not None else 25
    if lo > hi:
        lo, hi = hi, lo
    acc: set[str] = set()
    for age_int in range(lo, hi + 1):
        acc |= _age_year_to_buckets(age_int)
    order = ["1.5-3", "3-5", "5-7", "7+"]
    return [b for b in order if b in acc]


def _school_classes_to_age_bounds(class_from: int | None, class_to: int | None) -> tuple[int | None, int | None]:
    if class_from is None:
        return None, None
    cf = int(class_from)
    ct = int(class_to) if class_to is not None else cf
    if cf > ct:
        cf, ct = ct, cf
    age_min = 6 + cf
    age_max = 7 + ct
    return age_min, age_max


def _preference_ranking_inputs(pref: UserDistrictPreference) -> tuple[list[str], list[str], list[str], list[str], bool]:
    type_codes: list[str] = []
    if pref.preferred_institution_type_id:
        type_codes = [pref.preferred_institution_type_id]
    pt = pref.preferred_institution_type_id
    if pt in {"school", "school_internat"}:
        amin, amax = _school_classes_to_age_bounds(pref.school_class_from, pref.school_class_to)
        age_buckets = _years_span_to_age_buckets(amin, amax)
    else:
        age_buckets = _years_span_to_age_buckets(pref.child_age_min, pref.child_age_max)
    cond = pref.preferred_condition_codes if isinstance(pref.preferred_condition_codes, list) else []
    acc = pref.preferred_accessibility_codes if isinstance(pref.preferred_accessibility_codes, list) else []
    return type_codes, age_buckets, cond, acc, bool(pref.prefer_aoop)


def _search_match(inst_payload: dict, search_term: str) -> bool:
    if not search_term:
        return True
    q = _normalize_search_value(search_term)
    if not q:
        return True

    name = _normalize_search_value(inst_payload.get("name") or "")
    desc = _normalize_search_value(inst_payload.get("description") or "")
    director_name = _normalize_search_value(((inst_payload.get("director") or {}).get("name")) or "")
    website = _normalize_search_value(inst_payload.get("website") or "")
    inst_number_match = re.search(r"\d+", inst_payload.get("name") or "")
    query_number_match = re.search(r"\d+", search_term or "")
    inst_number = inst_number_match.group(0) if inst_number_match else None
    query_number = query_number_match.group(0) if query_number_match else None

    if q in name or q in desc or q in director_name or q in website:
        return True
    if query_number and inst_number and query_number == inst_number:
        return True

    synonyms = {
        "preschool": ["детсад", "детскийсад", "садик", "дс", "сад"],
        "school": ["школа", "шк", "сош", "лицей", "гимназия"],
        "school_internat": ["интернат", "школаинтернат"],
        "spo": ["спо", "техникум", "колледж", "училище"],
        "vo": ["во", "вуз", "университет", "институт", "академия"],
    }
    normalized_term = _normalize_search_value(search_term)
    for type_code, words in synonyms.items():
        if any(word in normalized_term for word in words) and inst_payload.get("type") == type_code:
            if not query_number:
                return True
            if inst_number and query_number == inst_number:
                return True
    return False


def _calc_relevance(
    inst_payload: dict,
    selected_types: list[str],
    selected_ages: list[str],
    selected_conditions: list[str],
    selected_accessibility: list[str],
    aoop_selected: bool,
) -> float:
    weights = {
        "type": 0.5,
        "conditions": 0.15,
        "age": 0.15,
        "aoop": 0.1,
        "accessibility": 0.1,
    }
    inst_conditions = inst_payload.get("conditions") or []
    inst_accessibility = inst_payload.get("accessibility_criteria") or []
    inst_aoop = inst_payload.get("aoop_programs") or []

    type_score = 1.0 if not selected_types else (1.0 if inst_payload.get("type") in selected_types else 0.0)
    conditions_score = (
        1.0
        if not selected_conditions
        else sum(1 for c in selected_conditions if c in inst_conditions) / len(selected_conditions)
    )
    age_score = (
        1.0
        if not selected_ages
        else sum(1 for age in selected_ages if _matches_age_bucket(inst_payload, age)) / len(selected_ages)
    )
    aoop_score = 1.0 if not aoop_selected else (1.0 if len(inst_aoop) > 0 else 0.0)
    accessibility_score = (
        1.0
        if not selected_accessibility
        else sum(1 for c in selected_accessibility if c in inst_accessibility) / len(selected_accessibility)
    )
    return (
        weights["type"] * type_score
        + weights["conditions"] * conditions_score
        + weights["age"] * age_score
        + weights["aoop"] * aoop_score
        + weights["accessibility"] * accessibility_score
    )


def _value_changed(old_value, new_value) -> bool:
    return (old_value or "") != (new_value or "")


def _build_institution_update_message(old_data: dict | None, new_data: dict | None, institution_name: str) -> str:
    old_payload = old_data or {}
    new_payload = new_data or {}
    changed_parts: list[str] = []

    if _value_changed(old_payload.get("name"), new_payload.get("name")):
        changed_parts.append("название")
    if _value_changed(old_payload.get("description"), new_payload.get("description")):
        changed_parts.append("описание")
    if _value_changed(old_payload.get("address"), new_payload.get("address")):
        changed_parts.append("адрес")
    if _value_changed(old_payload.get("website"), new_payload.get("website")):
        changed_parts.append("сайт")
    if _value_changed(old_payload.get("aoop_url"), new_payload.get("aoop_url")):
        changed_parts.append("ссылка на АООП")
    if _value_changed(old_payload.get("district_id"), new_payload.get("district_id")):
        changed_parts.append("округ")
    if _value_changed(old_payload.get("type"), new_payload.get("type")):
        changed_parts.append("тип учреждения")

    old_min = (old_payload.get("range") or {}).get("min") if isinstance(old_payload.get("range"), dict) else old_payload.get("range_min")
    old_max = (old_payload.get("range") or {}).get("max") if isinstance(old_payload.get("range"), dict) else old_payload.get("range_max")
    new_min = (new_payload.get("range") or {}).get("min") if isinstance(new_payload.get("range"), dict) else new_payload.get("range_min")
    new_max = (new_payload.get("range") or {}).get("max") if isinstance(new_payload.get("range"), dict) else new_payload.get("range_max")
    if _value_changed(old_min, new_min) or _value_changed(old_max, new_max):
        changed_parts.append("возрастной диапазон")

    old_director = old_payload.get("director") or {}
    new_director = new_payload.get("director") or {}
    if any(
        _value_changed(old_director.get(key), new_director.get(key))
        for key in ("name", "phone", "email")
    ):
        changed_parts.append("данные директора")

    old_conditions = old_payload.get("conditions") or []
    new_conditions = new_payload.get("conditions") or []
    if sorted(old_conditions) != sorted(new_conditions):
        changed_parts.append("условия обучения")

    old_admission = old_payload.get("conditionsAdmission") or []
    new_admission = new_payload.get("conditionsAdmission") or []
    if sorted(old_admission) != sorted(new_admission):
        changed_parts.append("формы поступления")

    old_accessibility = old_payload.get("accessibility_criteria") or []
    new_accessibility = new_payload.get("accessibility_criteria") or []
    if sorted(old_accessibility) != sorted(new_accessibility):
        changed_parts.append("критерии доступности")

    old_aoop = sorted((p.get("name") or "").strip() for p in (old_payload.get("aoop_programs") or []) if isinstance(p, dict))
    new_aoop = sorted((p.get("name") or "").strip() for p in (new_payload.get("aoop_programs") or []) if isinstance(p, dict))
    if old_aoop != new_aoop:
        changed_parts.append("программы АООП")

    if not changed_parts:
        return f'В учреждение "{institution_name}" внесены изменения.'
    return f'В учреждении "{institution_name}" обновлены: {", ".join(changed_parts)}.'


@require_GET
def get_districts(request: HttpRequest):
    districts = District.objects.order_by("name").values("id", "name")
    district_map = {row["id"]: row["name"] for row in districts}
    return JsonResponse({"districts": district_map})


@require_GET
def get_links(request: HttpRequest):
    district_id = request.GET.get("district_id")
    if not district_id:
        return JsonResponse({"error": "district_id required"}, status=400)
    links = Link.objects.filter(district_id=district_id).order_by("id").values(
        "linksTo", "linkToInstitution"
    )
    payload = [{"linksto": l["linksTo"], "linktoinstitution": l["linkToInstitution"]} for l in links]
    return JsonResponse({"links": payload})


@require_GET
def get_institutions(request: HttpRequest):
    district_id = request.GET.get("district_id")
    if not district_id:
        return JsonResponse({"error": "district_id required and must be integer"}, status=400)

    queryset = (
        Institution.objects.filter(district_id=district_id)
        .select_related("director", "type", "district")
        .prefetch_related("conditions", "admission", "aoop_programs", "accessibility_criteria")
        .order_by("name")
    )
    page, page_size = _get_pagination(request)
    total_items = queryset.count()
    pagination = _build_pagination(total_items=total_items, page=page, page_size=page_size)
    start_idx = (pagination["page"] - 1) * page_size
    end_idx = start_idx + page_size
    queryset = queryset[start_idx:end_idx]
    institution_ids = list(queryset.values_list("id", flat=True))
    review_aggs = _get_reviews_aggregates_for_institutions(institution_ids)
    approved_reviews_map = _get_approved_reviews_for_institutions(institution_ids)
    favorite_ids = _get_favorite_ids_for_request(request)
    return JsonResponse(
        {
            "institutions": [
                _serialize_institution(inst, favorite_ids, review_aggs, approved_reviews_map) for inst in queryset
            ],
            "pagination": pagination,
        }
    )


@require_GET
def search_institutions(request: HttpRequest):
    district_id = request.GET.get("district_id")
    if not district_id:
        return JsonResponse({"error": "district_id required"}, status=400)

    selected_types = [v for v in request.GET.getlist("type") if v]
    selected_ages = [v for v in request.GET.getlist("age") if v]
    selected_conditions = [v for v in request.GET.getlist("condition") if v]
    selected_accessibility = [v for v in request.GET.getlist("accessibility") if v]
    aoop_selected = request.GET.get("aoop") in {"1", "true", "yes"}
    search_term = (request.GET.get("q") or "").strip()

    queryset = (
        Institution.objects.filter(district_id=district_id)
        .select_related("director", "type", "district")
        .prefetch_related("conditions", "admission", "aoop_programs", "accessibility_criteria")
    )
    institution_ids = list(queryset.values_list("id", flat=True))
    review_aggs = _get_reviews_aggregates_for_institutions(institution_ids)
    approved_reviews_map = _get_approved_reviews_for_institutions(institution_ids)
    favorite_ids = _get_favorite_ids_for_request(request)
    institutions = [_serialize_institution(inst, favorite_ids, review_aggs, approved_reviews_map) for inst in queryset]
    institutions = [inst for inst in institutions if _search_match(inst, search_term)]
    ranked_with_relevance = [
        (
            inst,
            _calc_relevance(
                inst,
                selected_types,
                selected_ages,
                selected_conditions,
                selected_accessibility,
                aoop_selected,
            ),
        )
        for inst in institutions
    ]
    ranked_with_relevance.sort(key=lambda pair: pair[1], reverse=True)
    ranked = []
    for inst, relevance in ranked_with_relevance:
        inst_with_score = dict(inst)
        inst_with_score["relevance"] = round(float(relevance), 3)
        ranked.append(inst_with_score)
    page, page_size = _get_pagination(request)
    total_items = len(ranked)
    pagination = _build_pagination(total_items=total_items, page=page, page_size=page_size)
    start_idx = (pagination["page"] - 1) * page_size
    end_idx = start_idx + page_size
    return JsonResponse({"institutions": ranked[start_idx:end_idx], "pagination": pagination})


@require_GET
def get_institution(request: HttpRequest):
    inst_id = request.GET.get("id")
    if not inst_id:
        return JsonResponse({"error": "id required and must be integer"}, status=400)
    inst = get_object_or_404(
        Institution.objects.select_related("director", "type", "district").prefetch_related(
            "conditions", "admission", "aoop_programs", "accessibility_criteria"
        ),
        pk=inst_id,
    )
    review_aggs = _get_reviews_aggregates_for_institutions([inst.id])
    approved_reviews_map = _get_approved_reviews_for_institutions([inst.id])
    favorite_ids = _get_favorite_ids_for_request(request)
    return JsonResponse({"institution": _serialize_institution(inst, favorite_ids, review_aggs, approved_reviews_map)})


def _require_portal_user(request: HttpRequest):
    user = request.user
    if not user.is_authenticated:
        return None
    if user.is_superuser or user.groups.filter(name="administrators").exists():
        return None
    if _user_is_blocked(user):
        return None
    return user


@require_GET
def get_favorites(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"favorite_ids": []})
    favorite_ids = list(Favorite.objects.filter(user=user).values_list("institution_id", flat=True))
    return JsonResponse({"favorite_ids": favorite_ids})


@csrf_exempt
@require_POST
def update_profile(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)

    data = _parse_json(request)
    if data is None:
        return JsonResponse({"error": "Invalid JSON"}, status=400)

    first_name = (data.get("first_name") or "").strip()
    username = (data.get("username") or "").strip()
    email = (data.get("email") or "").strip().lower()
    notify_new_institutions = bool(data.get("notify_new_institutions", False))

    if not first_name:
        return JsonResponse({"error": "Имя не может быть пустым"}, status=400)
    if not username:
        return JsonResponse({"error": "Логин не может быть пустым"}, status=400)
    if not email:
        return JsonResponse({"error": "Почта не может быть пустой"}, status=400)

    User = get_user_model()
    username_taken = User.objects.filter(username=username).exclude(id=user.id).exists()
    if username_taken:
        return JsonResponse({"error": "Этот username уже занят"}, status=409)
    email_taken = User.objects.filter(email__iexact=email).exclude(id=user.id).exists()
    if email_taken:
        return JsonResponse({"error": "Эта почта уже используется"}, status=409)

    user.first_name = first_name
    user.username = username
    user.email = email
    user.save(update_fields=["first_name", "username", "email"])
    preference, _ = UserDistrictPreference.objects.get_or_create(user=user)
    pref_fields: list[str] = []
    if "district_id" in data:
        district_id = data.get("district_id")
        district_obj = None
        if district_id not in (None, "", 0, "0"):
            district_obj = District.objects.filter(id=district_id).first()
            if not district_obj:
                return JsonResponse({"error": "Выбранный район не найден"}, status=400)
            preference.district = district_obj
        else:
            preference.district = None
        pref_fields.extend(["district", "updated_at"])
    if "notify_new_institutions" in data:
        preference.notify_new_institutions = bool(
            notify_new_institutions and preference.district_id is not None
        )
        pref_fields.extend(["notify_new_institutions", "updated_at"])
    if pref_fields:
        preference.save(update_fields=list(dict.fromkeys(pref_fields)))
    return JsonResponse(
        {
            "success": True,
            "profile": {
                "first_name": user.first_name,
                "username": user.username,
                "email": user.email,
                "district_id": preference.district_id,
                "notify_new_institutions": preference.notify_new_institutions,
            },
        }
    )


def _coerce_str_code_list(raw) -> list[str]:
    if raw is None:
        return []
    if not isinstance(raw, list):
        return []
    return [str(item).strip() for item in raw if str(item).strip()]


def _parse_optional_positive_int(raw) -> int | None:
    if raw in (None, "", 0, "0"):
        return None
    try:
        v = int(raw)
    except (TypeError, ValueError):
        return None
    if v < 0:
        return None
    return v


@csrf_exempt
@require_POST
def update_portal_main_info(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    data = _parse_json(request)
    if data is None:
        return JsonResponse({"error": "Invalid JSON"}, status=400)

    pref, _ = UserDistrictPreference.objects.get_or_create(user=user)

    district_id = data.get("district_id")
    district_obj = None
    if district_id not in (None, "", 0, "0"):
        district_obj = District.objects.filter(id=district_id).first()
        if not district_obj:
            return JsonResponse({"error": "Выбранный район не найден"}, status=400)
    pref.district = district_obj

    type_code = data.get("preferred_institution_type")
    if type_code in (None, "", 0, "0"):
        pref.preferred_institution_type = None
    else:
        type_code = str(type_code).strip()
        if not InstitutionType.objects.filter(pk=type_code).exists():
            return JsonResponse({"error": "Неизвестный тип учреждения"}, status=400)
        pref.preferred_institution_type_id = type_code

    pt = pref.preferred_institution_type_id
    if pt in {"school", "school_internat"}:
        scf = _parse_optional_positive_int(data.get("school_class_from"))
        sct = _parse_optional_positive_int(data.get("school_class_to"))
        if scf is None:
            return JsonResponse({"error": "Укажите класс (от) для школы или школы-интерната"}, status=400)
        if scf < 1 or scf > 11:
            return JsonResponse({"error": "Класс должен быть от 1 до 11"}, status=400)
        if sct is not None and (sct < 1 or sct > 11):
            return JsonResponse({"error": "Класс должен быть от 1 до 11"}, status=400)
        pref.school_class_from = scf
        pref.school_class_to = sct if sct is not None else scf
        pref.child_age_min = None
        pref.child_age_max = None
    else:
        amin = _parse_optional_positive_int(data.get("child_age_min"))
        amax = _parse_optional_positive_int(data.get("child_age_max"))
        if amin is not None and amin > 120:
            return JsonResponse({"error": "Некорректный возраст"}, status=400)
        if amax is not None and amax > 120:
            return JsonResponse({"error": "Некорректный возраст"}, status=400)
        if amin is not None and amax is not None and amin > amax:
            return JsonResponse({"error": "Возраст «от» не может быть больше «до»"}, status=400)
        pref.child_age_min = amin
        pref.child_age_max = amax
        pref.school_class_from = None
        pref.school_class_to = None

    cond_codes = _coerce_str_code_list(data.get("preferred_condition_codes"))
    valid_conditions = set(ConditionType.objects.values_list("code", flat=True))
    if any(c not in valid_conditions for c in cond_codes):
        return JsonResponse({"error": "Неизвестный код условия обучения"}, status=400)
    pref.preferred_condition_codes = cond_codes

    acc_codes = _coerce_str_code_list(data.get("preferred_accessibility_codes"))
    valid_acc = set(AccessibilityCriterionType.objects.values_list("code", flat=True))
    if any(c not in valid_acc for c in acc_codes):
        return JsonResponse({"error": "Неизвестный код критерия доступности"}, status=400)
    pref.preferred_accessibility_codes = acc_codes

    pref.prefer_aoop = bool(data.get("prefer_aoop", False))

    pref.save(
        update_fields=[
            "district",
            "preferred_institution_type",
            "child_age_min",
            "child_age_max",
            "school_class_from",
            "school_class_to",
            "preferred_condition_codes",
            "preferred_accessibility_codes",
            "prefer_aoop",
            "updated_at",
        ]
    )
    return JsonResponse(
        {
            "success": True,
            "preference": {
                "district_id": pref.district_id,
                "preferred_institution_type": pref.preferred_institution_type_id,
                "child_age_min": pref.child_age_min,
                "child_age_max": pref.child_age_max,
                "school_class_from": pref.school_class_from,
                "school_class_to": pref.school_class_to,
                "preferred_condition_codes": list(pref.preferred_condition_codes or []),
                "preferred_accessibility_codes": list(pref.preferred_accessibility_codes or []),
                "prefer_aoop": pref.prefer_aoop,
            },
        }
    )


@require_GET
def profile_match_institutions(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    pref = UserDistrictPreference.objects.filter(user=user).first()
    if not pref or not pref.district_id:
        return JsonResponse(
            {"error": "Укажите и сохраните район в основной информации профиля.", "institutions": []},
            status=400,
        )
    cond_codes = pref.preferred_condition_codes if isinstance(pref.preferred_condition_codes, list) else []
    valid_conditions = set(ConditionType.objects.values_list("code", flat=True))
    if any(c not in valid_conditions for c in cond_codes):
        return JsonResponse({"error": "Сохранённые условия обучения недействительны.", "institutions": []}, status=400)

    qs = Institution.objects.filter(district_id=pref.district_id)
    for code in cond_codes:
        qs = qs.filter(conditions__code=code)
    qs = (
        qs.distinct()
        .select_related("director", "type", "district")
        .prefetch_related("conditions", "admission", "aoop_programs", "accessibility_criteria")
    )
    institution_ids = list(qs.values_list("id", flat=True))
    review_aggs = _get_reviews_aggregates_for_institutions(institution_ids)
    approved_reviews_map = _get_approved_reviews_for_institutions(institution_ids)
    favorite_ids = _get_favorite_ids_for_request(request)
    institutions = [_serialize_institution(inst, favorite_ids, review_aggs, approved_reviews_map) for inst in qs]
    type_codes, age_buckets, cond_rank, acc_rank, aoop_sel = _preference_ranking_inputs(pref)
    ranked_with_relevance = [
        (
            inst,
            _calc_relevance(
                inst,
                type_codes,
                age_buckets,
                cond_rank,
                acc_rank,
                aoop_sel,
            ),
        )
        for inst in institutions
    ]
    ranked_with_relevance.sort(key=lambda pair: pair[1], reverse=True)
    ranked = []
    for inst, relevance in ranked_with_relevance:
        inst_with_score = dict(inst)
        inst_with_score["relevance"] = round(float(relevance), 3)
        ranked.append(inst_with_score)
    page, page_size = _get_pagination(request)
    total_items = len(ranked)
    pagination = _build_pagination(total_items=total_items, page=page, page_size=page_size)
    start_idx = (pagination["page"] - 1) * page_size
    end_idx = start_idx + page_size
    return JsonResponse({"institutions": ranked[start_idx:end_idx], "pagination": pagination})


@csrf_exempt
@require_POST
def add_favorite(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    data = _parse_json(request) or {}
    institution_id = data.get("institution_id")
    if not institution_id:
        return JsonResponse({"error": "institution_id is required"}, status=400)
    institution = Institution.objects.filter(pk=institution_id).first()
    if not institution:
        return JsonResponse({"error": "Institution not found"}, status=404)
    Favorite.objects.get_or_create(user=user, institution=institution)
    return JsonResponse({"success": True, "is_favorite": True})


@csrf_exempt
@require_POST
def remove_favorite(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    data = _parse_json(request) or {}
    institution_id = data.get("institution_id")
    if not institution_id:
        return JsonResponse({"error": "institution_id is required"}, status=400)
    Favorite.objects.filter(user=user, institution_id=institution_id).delete()
    return JsonResponse({"success": True, "is_favorite": False})


@require_GET
def get_directors(request: HttpRequest):
    search = request.GET.get("search", "")
    qs = Director.objects.all()
    if search:
        qs = qs.filter(full_name__icontains=search)
    directors = list(qs.order_by("full_name").values("id", "full_name", "phone", "email"))
    return JsonResponse({"directors": directors})


@require_GET
def get_director(request: HttpRequest):
    director_id = request.GET.get("id")
    if not director_id:
        return JsonResponse({"error": "Director ID is required"}, status=400)
    director = get_object_or_404(Director, pk=director_id)
    return JsonResponse(
        {"director": {"id": director.id, "full_name": director.full_name, "phone": director.phone, "email": director.email}}
    )


@csrf_exempt
@require_POST
def login(request: HttpRequest):
    data = _parse_json(request)
    if not data or not data.get("login") or not data.get("password"):
        return JsonResponse({"error": "Invalid JSON or missing credentials"}, status=400)

    user = _authenticate_with_login_or_email(request, data["login"], data["password"])
    if not user:
        return JsonResponse({"error": "Invalid credentials"}, status=401)
    if _user_is_blocked(user):
        return _blocked_login_response(user)
    is_admin = user.is_superuser or user.groups.filter(name="administrators").exists()
    if not is_admin:
        return JsonResponse({"error": "Admin role is required"}, status=403)

    auth_login(request, user)
    return JsonResponse({"success": True})


@csrf_exempt
@require_POST
def logout(request: HttpRequest):
    auth_logout(request)
    return JsonResponse({"success": True})


@require_GET
def check_session(request: HttpRequest):
    user = request.user
    is_authenticated = bool(user and user.is_authenticated)
    is_admin = is_authenticated and (user.is_superuser or user.groups.filter(name="administrators").exists())
    account_blocked = False
    block_reason = ""
    if is_authenticated and not is_admin and _user_is_blocked(user):
        account_blocked = True
        block_reason = _user_block_reason(user)
        auth_logout(request)
        is_authenticated = False
        is_admin = False
    payload = {
        "loggedIn": bool(is_admin),
        "adminLoggedIn": bool(is_admin),
        "portalLoggedIn": bool(is_authenticated and not is_admin),
        "accountBlocked": account_blocked,
        "blockReason": block_reason,
    }
    if is_authenticated and not is_admin:
        payload["portalUserName"] = user.get_full_name() or user.username
    return JsonResponse(payload)


@csrf_exempt
@require_POST
def register(request: HttpRequest):
    data = _parse_json(request)
    required = ("login", "email", "password", "full_name")
    if not data or any(not data.get(field) for field in required):
        return JsonResponse({"error": "Invalid JSON or missing fields: login, email, password, full_name"}, status=400)

    User = get_user_model()
    email_s = data["email"].strip()
    login_s = email_s
    if User.objects.filter(Q(username=login_s) | Q(email=email_s)).exists():
        return JsonResponse({"error": "Login or email already exists"}, status=409)

    admin_group, _ = Group.objects.get_or_create(name="administrators")
    user = _create_user_with_repaired_sequence(
        username=login_s,
        email=email_s,
        password=data["password"],
        first_name=data["full_name"].strip(),
    )
    user.groups.add(admin_group)
    user.is_staff = True
    user.last_login = timezone.now()
    user.save(update_fields=["is_staff", "last_login"])
    auth_login(request, user)
    return JsonResponse({"success": True, "id": user.id}, status=201)


@csrf_exempt
@require_POST
def register_portal_user(request: HttpRequest):
    data = _parse_json(request)
    required = ("email", "password", "full_name")
    if not data or any(not data.get(field) for field in required):
        return JsonResponse({"error": "Invalid JSON or missing fields: email, password, full_name"}, status=400)

    User = get_user_model()
    email_s = data["email"].strip()
    login_s = email_s
    if User.objects.filter(Q(username=login_s) | Q(email=email_s)).exists():
        return JsonResponse({"error": "Login or email already exists"}, status=409)

    portal_group, _ = Group.objects.get_or_create(name="portal_users")
    user = _create_user_with_repaired_sequence(
        username=login_s,
        email=email_s,
        password=data["password"],
        first_name=data["full_name"].strip(),
    )
    user.groups.add(portal_group)
    user.last_login = timezone.now()
    user.save(update_fields=["last_login"])
    auth_login(request, user)
    return JsonResponse({"success": True, "id": user.id}, status=201)


@csrf_exempt
@require_POST
def login_portal_user(request: HttpRequest):
    data = _parse_json(request)
    if not data or not data.get("login") or not data.get("password"):
        return JsonResponse({"error": "Invalid JSON or missing credentials"}, status=400)

    user = _authenticate_with_login_or_email(request, data["login"], data["password"])
    if not user:
        return JsonResponse({"error": "Invalid credentials"}, status=401)
    if _user_is_blocked(user):
        return _blocked_login_response(user)
    is_admin = user.is_superuser or user.groups.filter(name="administrators").exists()
    if is_admin:
        return JsonResponse({"error": "Portal user role is required"}, status=403)

    auth_login(request, user)
    return JsonResponse({"success": True})


def _upsert_director(data):
    director_data = data.get("director") or {}
    if not director_data.get("name"):
        return None
    director_id = director_data.get("id")
    if director_id:
        director = Director.objects.filter(id=director_id).first()
        if director:
            director.full_name = director_data.get("name")
            director.phone = director_data.get("phone")
            director.email = director_data.get("email")
            director.save()
            return director
    return Director.objects.create(
        full_name=director_data.get("name"),
        phone=director_data.get("phone"),
        email=director_data.get("email"),
    )


def _sync_relations(inst: Institution, data):
    conditions = data.get("conditions", [])
    admissions = data.get("conditionsAdmission", [])
    accessibility_criteria = data.get("accessibility_criteria", [])
    inst.conditions.set(ConditionType.objects.filter(code__in=conditions))
    inst.admission.set(AdmissionType.objects.filter(code__in=admissions))
    inst.accessibility_criteria.set(
        AccessibilityCriterionType.objects.filter(code__in=accessibility_criteria)
    )
    inst.aoop_programs.all().delete()
    AoopProgram.objects.bulk_create(
        [AoopProgram(institution=inst, name=p.get("name", "")) for p in data.get("aoop_programs", []) if p.get("name")]
    )


def _require_admin(request: HttpRequest):
    user = request.user
    if not user.is_authenticated:
        return None
    if user.is_superuser or user.groups.filter(name="administrators").exists():
        return user
    return None


@csrf_exempt
@require_GET
def export_admin_report(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)

    export_format = (request.GET.get("format") or "pdf").strip().lower()
    if export_format not in {"pdf", "excel", "csv"}:
        return JsonResponse({"error": "Формат не поддерживается"}, status=400)

    now = timezone.localtime()
    month_value = (request.GET.get("month") or "").strip()
    if month_value:
        try:
            requested_month = datetime.strptime(month_value, "%Y-%m")
            period_start = timezone.make_aware(datetime(requested_month.year, requested_month.month, 1))
        except ValueError:
            return JsonResponse({"error": "Некорректный месяц. Используйте формат YYYY-MM"}, status=400)
    else:
        period_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    if period_start.month == 12:
        period_end = period_start.replace(year=period_start.year + 1, month=1, day=1)
    else:
        period_end = period_start.replace(month=period_start.month + 1, day=1)

    total_institutions = Institution.objects.count()
    institutions_month_delta = Institution.objects.filter(created_at__gte=period_start, created_at__lt=period_end).count()
    User = get_user_model()
    total_users = User.objects.count()
    users_month_delta = User.objects.filter(date_joined__gte=period_start, date_joined__lt=period_end).count()
    moderation_total = InstitutionReview.objects.filter(
        status__in=[InstitutionReview.STATUS_PENDING, InstitutionReview.STATUS_DISAPPROVED, "rejected"]
    ).count()
    moderation_month_delta = InstitutionReview.objects.filter(
        status__in=[InstitutionReview.STATUS_PENDING, InstitutionReview.STATUS_DISAPPROVED, "rejected"],
        created_at__gte=period_start,
        created_at__lt=period_end,
    ).count()

    approved_reviews_qs = InstitutionReview.objects.filter(status=InstitutionReview.STATUS_APPROVED)
    approved_reviews_total = approved_reviews_qs.count()
    avg_approved_rating = approved_reviews_qs.aggregate(value=Avg("rating")).get("value")
    avg_approved_rating_value = round(float(avg_approved_rating), 2) if avg_approved_rating is not None else 0.0

    distribution_rows = list(
        Institution.objects.values("type__name_ru")
        .annotate(total=Count("id"))
        .order_by("-total")
    )
    top_type_row = distribution_rows[0] if distribution_rows else None
    top_type_name = (top_type_row or {}).get("type__name_ru") or "n/a"
    top_type_count = int((top_type_row or {}).get("total") or 0)
    type_distribution_total = sum(int(row.get("total") or 0) for row in distribution_rows) or 1
    type_distribution_lines = [
        (
            (row.get("type__name_ru") or "Без типа"),
            int(row.get("total") or 0),
            (int(row.get("total") or 0) * 100.0) / type_distribution_total,
        )
        for row in distribution_rows
    ]

    top_user_row = (
        User.objects.filter(is_superuser=False)
        .annotate(reviews_count=Count("institution_reviews"))
        .order_by("-reviews_count", "username")
        .values("first_name", "last_name", "username", "reviews_count")
        .first()
    )
    top_user_name = "n/a"
    top_user_reviews = 0
    if top_user_row:
        full_name = " ".join(
            part for part in [top_user_row.get("first_name"), top_user_row.get("last_name")] if part
        ).strip()
        top_user_name = full_name or top_user_row.get("username") or "n/a"
        top_user_reviews = int(top_user_row.get("reviews_count") or 0)

    users_list_rows = (
        User.objects.annotate(reviews_count=Count("institution_reviews"))
        .order_by("-reviews_count", "username", "id")
        .values("first_name", "last_name", "username", "email", "reviews_count")
    )
    users_list = []
    for row in users_list_rows:
        full_name = " ".join(part for part in [row["first_name"], row["last_name"]] if part).strip()
        users_list.append(
            {
                "display_name": full_name or row["username"] or (row["email"] or "Пользователь"),
                "username": row["username"] or "",
                "reviews_count": int(row["reviews_count"] or 0),
            }
        )

    global_rating_mean = _get_global_approved_rating_mean()
    institutions_list_rows = (
        Institution.objects.annotate(
            avg_rating=Avg("reviews__rating", filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED)),
            approved_reviews_count=Count("reviews", filter=Q(reviews__status=InstitutionReview.STATUS_APPROVED)),
        )
        .select_related("district")
        .order_by("-avg_rating", "name")
        .values("name", "district__name", "avg_rating", "approved_reviews_count")
    )
    institutions_list = [
        {
            "name": row["name"] or "Учреждение",
            "district_name": row["district__name"] or "",
            "avg_rating": _calc_bayesian_rating(
                float(row["avg_rating"]) if row["avg_rating"] is not None else None,
                int(row.get("approved_reviews_count") or 0),
                global_rating_mean,
            ),
        }
        for row in institutions_list_rows
    ]

    if export_format == "csv":
        from io import StringIO

        csv_buffer = StringIO()
        writer = csv.writer(csv_buffer, delimiter=";")
        writer.writerow(["Admin Dashboard Report"])
        writer.writerow(["Generated", now.strftime("%Y-%m-%d %H:%M")])
        writer.writerow(
            [
                "Selected period",
                period_start.strftime("%Y-%m-%d"),
                (period_end - timedelta(days=1)).strftime("%Y-%m-%d"),
            ]
        )
        writer.writerow([])
        writer.writerow(["Core metrics"])
        writer.writerow(["Institutions total", total_institutions, "New in selected month", institutions_month_delta])
        writer.writerow(["Registered users", total_users, "New in selected month", users_month_delta])
        writer.writerow(["Reviews on moderation", moderation_total, "New in selected month", moderation_month_delta])
        writer.writerow(["Approved reviews total", approved_reviews_total])
        writer.writerow(["Average approved institution rating", f"{avg_approved_rating_value:.2f}"])
        writer.writerow([])
        writer.writerow(["Additional insights"])
        writer.writerow(["Largest institution type", top_type_name, top_type_count])
        writer.writerow(["Most active reviewer", top_user_name, top_user_reviews])
        writer.writerow([])
        writer.writerow(["Distribution by institution type"])
        writer.writerow(["#", "Type", "Count", "Percent"])
        for idx, item in enumerate(type_distribution_lines, start=1):
            writer.writerow([idx, item[0], item[1], f"{item[2]:.2f}%"])
        writer.writerow([])
        writer.writerow(["All users and review counts"])
        writer.writerow(["#", "Display name", "Username", "Reviews"])
        for idx, item in enumerate(users_list, start=1):
            writer.writerow([idx, item["display_name"], item["username"], item["reviews_count"]])
        writer.writerow([])
        writer.writerow(["All institutions and average ratings"])
        writer.writerow(["#", "Institution", "District", "Average rating"])
        for idx, item in enumerate(institutions_list, start=1):
            writer.writerow(
                [
                    idx,
                    item["name"],
                    item["district_name"],
                    "n/a" if item["avg_rating"] is None else f"{item['avg_rating']:.2f}",
                ]
            )

        csv_content = csv_buffer.getvalue()
        filename = f"admin_dashboard_report_{period_start.strftime('%Y%m')}_{now.strftime('%d_%H%M')}.csv"
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response.write("\ufeff")
        response.write(csv_content)
        return response

    if export_format == "excel":
        try:
            from io import BytesIO

            from openpyxl import Workbook
        except Exception:
            return JsonResponse({"error": "Excel библиотека не установлена. Установите openpyxl."}, status=500)

        workbook = Workbook()
        summary_ws = workbook.active
        summary_ws.title = "Summary"
        summary_ws.append(["Admin Dashboard Report"])
        summary_ws.append(["Generated", now.strftime("%Y-%m-%d %H:%M")])
        summary_ws.append(
            [
                "Selected period",
                period_start.strftime("%Y-%m-%d"),
                (period_end - timedelta(days=1)).strftime("%Y-%m-%d"),
            ]
        )
        summary_ws.append([])
        summary_ws.append(["Core metrics"])
        summary_ws.append(["Institutions total", total_institutions, "New in selected month", institutions_month_delta])
        summary_ws.append(["Registered users", total_users, "New in selected month", users_month_delta])
        summary_ws.append(["Reviews on moderation", moderation_total, "New in selected month", moderation_month_delta])
        summary_ws.append(["Approved reviews total", approved_reviews_total])
        summary_ws.append(["Average approved institution rating", avg_approved_rating_value])
        summary_ws.append([])
        summary_ws.append(["Additional insights"])
        summary_ws.append(["Largest institution type", top_type_name, top_type_count])
        summary_ws.append(["Most active reviewer", top_user_name, top_user_reviews])

        type_ws = workbook.create_sheet("Types")
        type_ws.append(["#", "Type", "Count", "Percent"])
        for idx, item in enumerate(type_distribution_lines, start=1):
            type_ws.append([idx, item[0], item[1], round(item[2], 2)])

        users_ws = workbook.create_sheet("Users")
        users_ws.append(["#", "Display name", "Username", "Reviews"])
        for idx, item in enumerate(users_list, start=1):
            users_ws.append([idx, item["display_name"], item["username"], item["reviews_count"]])

        institutions_ws = workbook.create_sheet("Institutions")
        institutions_ws.append(["#", "Institution", "District", "Average rating"])
        for idx, item in enumerate(institutions_list, start=1):
            institutions_ws.append(
                [idx, item["name"], item["district_name"], None if item["avg_rating"] is None else item["avg_rating"]]
            )

        excel_buffer = BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)
        filename = f"admin_dashboard_report_{period_start.strftime('%Y%m')}_{now.strftime('%d_%H%M')}.xlsx"
        response = HttpResponse(
            excel_buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    try:
        from io import BytesIO

        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except Exception:
        return JsonResponse({"error": "PDF библиотека не установлена. Установите reportlab."}, status=500)

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    font_name = "Helvetica"
    for candidate in [
        ("Arial", "C:/Windows/Fonts/arial.ttf"),
        ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ("DejaVuSans", "/usr/share/fonts/dejavu/DejaVuSans.ttf"),
        ("LiberationSans", "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        ("NotoSans", "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
    ]:
        try:
            pdfmetrics.registerFont(TTFont(candidate[0], candidate[1]))
            font_name = candidate[0]
            break
        except Exception:
            continue

    y = height - 48
    bottom_margin = 48

    def ensure_space(required_height: int = 14):
        nonlocal y
        if y - required_height < bottom_margin:
            pdf.showPage()
            y = height - 48

    def write_line(text: str, x: int = 48, step: int = 14, size: int = 10):
        nonlocal y
        ensure_space(step)
        pdf.setFont(font_name, size)
        pdf.drawString(x, y, text)
        y -= step

    def write_section_title(text: str):
        nonlocal y
        ensure_space(20)
        pdf.setFont(font_name, 12)
        pdf.drawString(40, y, text)
        y -= 16

    pdf.setTitle("Admin Dashboard Report")
    pdf.setFont(font_name, 16)
    pdf.drawString(40, y, "Admin Dashboard Report")
    y -= 22
    write_line(f"Generated: {now.strftime('%Y-%m-%d %H:%M')}", x=40, step=18, size=10)
    write_line(
        f"Selected period: {period_start.strftime('%Y-%m-%d')} to {(period_end - timedelta(days=1)).strftime('%Y-%m-%d')}",
        x=40,
        step=18,
        size=10,
    )
    y -= 8

    write_section_title("Core metrics")
    metrics_lines = [
        f"- Institutions total: {total_institutions} (new in selected month: +{institutions_month_delta})",
        f"- Registered users: {total_users} (new in selected month: +{users_month_delta})",
        f"- Reviews on moderation: {moderation_total} (new in selected month: +{moderation_month_delta})",
        f"- Approved reviews total: {approved_reviews_total}",
        f"- Average approved institution rating: {avg_approved_rating_value:.2f}",
    ]
    for line in metrics_lines:
        write_line(line)

    y -= 8
    write_section_title("Additional insights")
    write_line(f"- Largest institution type: {top_type_name} ({top_type_count})")
    write_line(f"- Most active reviewer: {top_user_name} ({top_user_reviews} reviews)")

    y -= 8
    write_section_title("Distribution by institution type")
    if type_distribution_lines:
        for idx, item in enumerate(type_distribution_lines, start=1):
            write_line(f"{idx}. {item[0]} - {item[1]} ({item[2]:.2f}%)")
    else:
        write_line("No institution types found.")

    y -= 8
    write_section_title("All users and review counts")
    if users_list:
        for idx, item in enumerate(users_list, start=1):
            username_part = f" (@{item['username']})" if item["username"] else ""
            write_line(f"{idx}. {item['display_name']}{username_part} - reviews: {item['reviews_count']}")
    else:
        write_line("No users found.")

    y -= 8
    write_section_title("All institutions and average ratings")
    if institutions_list:
        for idx, item in enumerate(institutions_list, start=1):
            district_part = f" [{item['district_name']}]" if item["district_name"] else ""
            rating_value = "n/a" if item["avg_rating"] is None else f"{item['avg_rating']:.2f}"
            write_line(f"{idx}. {item['name']}{district_part} - average rating: {rating_value}")
    else:
        write_line("No institutions found.")

    pdf.save()
    buffer.seek(0)

    filename = f"admin_dashboard_report_{period_start.strftime('%Y%m')}_{now.strftime('%d_%H%M')}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@csrf_exempt
@require_POST
def delete_user(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    user_id = data.get("user_id")
    if not user_id:
        return JsonResponse({"error": "user_id is required"}, status=400)
    User = get_user_model()
    target = User.objects.filter(id=user_id).first()
    if not target:
        return JsonResponse({"error": "User not found"}, status=404)
    if target.id == admin.id:
        return JsonResponse({"error": "Нельзя удалить текущего администратора"}, status=400)
    if target.is_superuser:
        return JsonResponse({"error": "Нельзя удалить суперпользователя"}, status=400)
    target.delete()
    return JsonResponse({"success": True})


@csrf_exempt
@require_POST
def set_user_block(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    user_id = data.get("user_id")
    block = bool(data.get("block"))
    reason = (data.get("reason") or "").strip()
    if not user_id:
        return JsonResponse({"error": "user_id is required"}, status=400)
    User = get_user_model()
    target = User.objects.filter(id=user_id).first()
    if not target:
        return JsonResponse({"error": "User not found"}, status=404)
    if target.id == admin.id:
        return JsonResponse({"error": "Нельзя изменить блокировку для собственной учётной записи"}, status=400)
    if target.is_superuser:
        return JsonResponse({"error": "Нельзя блокировать суперпользователя"}, status=400)
    state, _ = UserBlockState.objects.get_or_create(user=target)
    state.is_blocked = block
    state.block_reason = reason if block else ""
    state.blocked_at = timezone.now() if block else None
    state.save(update_fields=["is_blocked", "block_reason", "blocked_at"])
    return JsonResponse(
        {
            "success": True,
            "is_blocked": state.is_blocked,
            "block_reason": (state.block_reason or "").strip(),
        }
    )


@csrf_exempt
@require_POST
def create_review(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Только зарегистрированный пользователь может оставить отзыв"}, status=403)

    data = _parse_json(request)
    if data is None:
        return JsonResponse({"error": "Invalid JSON"}, status=400)

    institution_id = data.get("institution_id")
    comment = (data.get("comment") or "").strip()
    institution_rating = data.get("institution_rating")
    criteria_ratings = data.get("criteria_ratings")

    if not institution_id:
        return JsonResponse({"error": "institution_id is required"}, status=400)
    institution = Institution.objects.filter(pk=institution_id).first()
    if not institution:
        return JsonResponse({"error": "Учреждение не найдено"}, status=404)
    existing_review = InstitutionReview.objects.filter(user=user, institution=institution).exists()
    if existing_review:
        return JsonResponse({"error": "Вы уже оставили отзыв для этого учреждения"}, status=409)
    if criteria_ratings is None:
        criteria_ratings = {}
    if not isinstance(criteria_ratings, dict):
        return JsonResponse({"error": "criteria_ratings must be an object"}, status=400)
    try:
        institution_rating_value = int(institution_rating)
    except (TypeError, ValueError):
        return JsonResponse({"error": "Оценка учреждения должна быть целым числом от 1 до 5"}, status=400)
    if institution_rating_value < 1 or institution_rating_value > 5:
        return JsonResponse({"error": "Оценка учреждения должна быть от 1 до 5"}, status=400)

    institution_criteria_codes = set(
        institution.accessibility_criteria.values_list("code", flat=True)
    )
    normalized_scores: dict[str, int] = {}
    for code, score in criteria_ratings.items():
        if str(code) not in institution_criteria_codes:
            return JsonResponse({"error": f"Критерий {code} не относится к этому учреждению"}, status=400)
        try:
            rating_value = int(score)
        except (TypeError, ValueError):
            return JsonResponse({"error": f"Некорректная оценка для критерия {code}"}, status=400)
        if rating_value < 1 or rating_value > 5:
            return JsonResponse({"error": f"Оценка для критерия {code} должна быть от 1 до 5"}, status=400)
        normalized_scores[str(code)] = rating_value

    if institution_criteria_codes and set(normalized_scores.keys()) != institution_criteria_codes:
        return JsonResponse({"error": "Оцените все критерии доступности этого учреждения"}, status=400)

    review = InstitutionReview.objects.create(
        user=user,
        institution=institution,
        rating=institution_rating_value,
        criteria_scores=normalized_scores,
        comment=comment,
        status=InstitutionReview.STATUS_PENDING,
    )
    return JsonResponse(
        {
            "success": True,
            "review_id": review.id,
            "status": review.status,
            "rating": float(review.rating),
        },
        status=201,
    )


def _notify_review_author_status_changed(
    administrator,
    review: InstitutionReview,
    old_status: str,
    new_status: str,
) -> None:
    if old_status == new_status:
        return
    institution = review.institution
    if institution is None:
        return
    author = review.user
    inst_name = institution.name or "учреждении"
    messages = {
        InstitutionReview.STATUS_APPROVED: f"Ваш отзыв об учреждении «{inst_name}» одобрен и опубликован на карте.",
        InstitutionReview.STATUS_DISAPPROVED: f"Ваш отзыв об учреждении «{inst_name}» отклонён.",
        InstitutionReview.STATUS_PENDING: f"Ваш отзыв об учреждении «{inst_name}» снова отправлен на модерацию.",
    }
    message = messages.get(new_status) or f"Статус вашего отзыва об учреждении «{inst_name}» обновлён."
    with transaction.atomic():
        action_log = ActionLog.objects.create(
            administrator=administrator,
            action="UPDATE",
            entity="reviews",
            record_id=review.id,
            old_data={"status": old_status},
            new_data={"status": new_status},
        )
        ProfileNotification.objects.create(
            user=author,
            institution=institution,
            action_log=action_log,
            message=message,
        )


@csrf_exempt
@require_POST
def approve_review(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.select_related("institution", "user").filter(id=review_id).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)
    old_status = review.status
    review.status = InstitutionReview.STATUS_APPROVED
    review.save(update_fields=["status", "updated_at"])
    _notify_review_author_status_changed(admin, review, old_status, review.status)
    return JsonResponse({"success": True, "status": review.status})


@csrf_exempt
@require_POST
def disapprove_review(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.select_related("institution", "user").filter(id=review_id).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)
    old_status = review.status
    review.status = InstitutionReview.STATUS_DISAPPROVED
    review.save(update_fields=["status", "updated_at"])
    _notify_review_author_status_changed(admin, review, old_status, review.status)
    return JsonResponse({"success": True, "status": review.status})


@csrf_exempt
@require_POST
def edit_review_placeholder(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.select_related("institution", "user").filter(id=review_id).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)

    raw_status = (data.get("status") or "").strip()
    status_value = "disapproved" if raw_status == "rejected" else raw_status
    allowed_statuses = {
        InstitutionReview.STATUS_PENDING,
        InstitutionReview.STATUS_APPROVED,
        InstitutionReview.STATUS_DISAPPROVED,
    }
    if status_value not in allowed_statuses:
        return JsonResponse({"error": "Некорректный статус модерации"}, status=400)

    comment = (data.get("comment") or "").strip()
    old_status = review.status
    review.status = status_value
    review.comment = comment
    review.save(update_fields=["status", "comment", "updated_at"])
    _notify_review_author_status_changed(admin, review, old_status, status_value)
    return JsonResponse({"success": True, "status": review.status, "comment": review.comment})


@require_GET
def get_admin_institution_reviews(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    inst_id = request.GET.get("institution_id")
    if not inst_id or not str(inst_id).isdigit():
        return JsonResponse({"error": "institution_id required"}, status=400)
    institution = Institution.objects.filter(pk=int(inst_id)).first()
    if not institution:
        return JsonResponse({"error": "Учреждение не найдено"}, status=404)
    reviews_qs = (
        InstitutionReview.objects.filter(institution=institution)
        .select_related("user")
        .order_by("-created_at", "-id")
    )
    reviews = []
    for rev in reviews_qs:
        user_name = rev.user.get_full_name() or rev.user.username or "Пользователь"
        reviews.append(
            {
                "id": rev.id,
                "user_name": user_name,
                "created_at_display": rev.created_at.strftime("%d.%m.%Y %H:%M"),
                "rating": float(rev.rating or 0),
                "comment": rev.comment or "",
                "status": rev.status,
            }
        )
    return JsonResponse(
        {
            "institution": {
                "id": institution.id,
                "name": institution.name or "",
                "district_name": institution.district.name if institution.district_id else "Район не указан",
                "type_name": institution.type.name_ru if institution.type_id else "Без типа",
            },
            "reviews": reviews,
        }
    )


@csrf_exempt
@require_POST
def delete_review_by_admin(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.filter(id=review_id).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)
    review.delete()
    return JsonResponse({"success": True})


@csrf_exempt
@require_GET
def get_my_review(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    review_id = request.GET.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = (
        InstitutionReview.objects.filter(id=review_id, user=user)
        .select_related("institution__district", "institution__type")
        .first()
    )
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)

    institution = review.institution
    criteria_scores = review.criteria_scores or {}
    criteria_rows = []
    if institution:
        criteria_qs = institution.accessibility_criteria.order_by("name_ru").values("code", "name_ru")
        for criterion in criteria_qs:
            score_value = criteria_scores.get(criterion["code"])
            try:
                parsed_score = int(score_value)
            except (TypeError, ValueError):
                parsed_score = None
            criteria_rows.append(
                {
                    "code": criterion["code"],
                    "name": criterion["name_ru"],
                    "score": parsed_score,
                }
            )

    status_label_map = {
        InstitutionReview.STATUS_PENDING: "На модерации",
        InstitutionReview.STATUS_APPROVED: "Опубликован",
        InstitutionReview.STATUS_DISAPPROVED: "Отклонен",
        "rejected": "Отклонен",
    }
    return JsonResponse(
        {
            "review": {
                "id": review.id,
                "institution_name": institution.name if institution else "Учреждение",
                "district_name": institution.district.name if institution and institution.district_id else "Город не указан",
                "type_name": institution.type.name_ru if institution and institution.type_id else "Тип не указан",
                "created_at": review.created_at.strftime("%d.%m.%Y %H:%M"),
                "status_value": review.status,
                "status_label": status_label_map.get(review.status, "На модерации"),
                "institution_rating": float(review.rating or 0),
                "comment": review.comment or "",
                "criteria": criteria_rows,
            }
        }
    )


@csrf_exempt
@require_POST
def update_my_review(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.filter(id=review_id, user=user).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)

    comment = (data.get("comment") or "").strip()
    institution_rating = data.get("institution_rating")
    criteria_ratings = data.get("criteria_ratings")
    if criteria_ratings is None:
        criteria_ratings = {}
    if not isinstance(criteria_ratings, dict):
        return JsonResponse({"error": "criteria_ratings must be an object"}, status=400)
    try:
        institution_rating_value = int(institution_rating)
    except (TypeError, ValueError):
        return JsonResponse({"error": "Оценка учреждения должна быть целым числом от 1 до 5"}, status=400)
    if institution_rating_value < 1 or institution_rating_value > 5:
        return JsonResponse({"error": "Оценка учреждения должна быть от 1 до 5"}, status=400)

    institution_criteria_codes = set(
        review.institution.accessibility_criteria.values_list("code", flat=True)
    )
    normalized_scores: dict[str, int] = {}
    for code, score in criteria_ratings.items():
        if str(code) not in institution_criteria_codes:
            return JsonResponse({"error": f"Критерий {code} не относится к этому учреждению"}, status=400)
        try:
            rating_value = int(score)
        except (TypeError, ValueError):
            return JsonResponse({"error": f"Некорректная оценка для критерия {code}"}, status=400)
        if rating_value < 1 or rating_value > 5:
            return JsonResponse({"error": f"Оценка для критерия {code} должна быть от 1 до 5"}, status=400)
        normalized_scores[str(code)] = rating_value

    if institution_criteria_codes and set(normalized_scores.keys()) != institution_criteria_codes:
        return JsonResponse({"error": "Оцените все критерии доступности этого учреждения"}, status=400)

    review.rating = institution_rating_value
    review.criteria_scores = normalized_scores
    review.comment = comment
    review.status = InstitutionReview.STATUS_PENDING
    review.save(update_fields=["rating", "criteria_scores", "comment", "status", "updated_at"])
    return JsonResponse(
        {
            "success": True,
            "status": review.status,
            "comment": review.comment,
            "institution_rating": float(review.rating or 0),
        }
    )


@csrf_exempt
@require_POST
def delete_my_review(request: HttpRequest):
    user = _require_portal_user(request)
    if not user:
        return JsonResponse({"error": "Portal user role is required"}, status=403)
    data = _parse_json(request) or {}
    review_id = data.get("review_id")
    if not review_id:
        return JsonResponse({"error": "review_id is required"}, status=400)
    review = InstitutionReview.objects.filter(id=review_id, user=user).first()
    if not review:
        return JsonResponse({"error": "Отзыв не найден"}, status=404)
    review.delete()
    return JsonResponse({"success": True})


@csrf_exempt
@require_POST
@transaction.atomic
def create_institution(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request)
    if not data:
        return JsonResponse({"error": "Invalid JSON"}, status=400)
    name = (data.get("name") or "").strip()
    if not name:
        return JsonResponse({"error": "Название учреждения обязательно"}, status=400)
    if Institution.objects.filter(name=name).exists():
        return JsonResponse({"error": "Учреждение с таким названием уже существует"}, status=409)

    director = _upsert_director(data)
    inst = Institution.objects.create(
        name=name,
        district_id=data.get("district_id"),
        type_id=data.get("type"),
        director=director,
        description=data.get("description"),
        range_min=(data.get("range") or {}).get("min"),
        range_max=(data.get("range") or {}).get("max"),
        website=data.get("website"),
        aoop_url=data.get("aoop_url"),
        address=data.get("address"),
    )
    _sync_relations(inst, data)
    action_log = ActionLog.objects.create(
        administrator=admin,
        action="CREATE",
        entity="institutions",
        record_id=inst.id,
        new_data=data,
    )
    subscribed_user_ids = list(
        UserDistrictPreference.objects.filter(
            district_id=inst.district_id,
            notify_new_institutions=True,
            user__is_superuser=False,
        )
        .exclude(user__groups__name="administrators")
        .values_list("user_id", flat=True)
        .distinct()
    )
    if subscribed_user_ids:
        district_name = inst.district.name if inst.district_id else "указанном районе"
        message = f'В вашем районе "{district_name}" добавлено новое учреждение "{inst.name}".'
        ProfileNotification.objects.bulk_create(
            [
                ProfileNotification(
                    user_id=user_id,
                    institution=inst,
                    action_log=action_log,
                    message=message,
                )
                for user_id in subscribed_user_ids
            ],
            ignore_conflicts=True,
        )
    return JsonResponse({"id": inst.id}, status=201)


@csrf_exempt
@require_POST
@transaction.atomic
def update_institution(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request)
    if not data or not data.get("id"):
        return JsonResponse({"error": "Invalid JSON or missing id"}, status=400)
    name = (data.get("name") or "").strip()
    if not name:
        return JsonResponse({"error": "Название учреждения обязательно"}, status=400)

    inst = get_object_or_404(Institution, id=data["id"])
    if Institution.objects.filter(name=name).exclude(id=inst.id).exists():
        return JsonResponse({"error": "Учреждение с таким названием уже существует"}, status=409)
    old_data = _serialize_institution(inst)
    director = _upsert_director(data)

    inst.name = name
    inst.district_id = data.get("district_id")
    inst.type_id = data.get("type")
    inst.director = director
    inst.description = data.get("description")
    inst.range_min = (data.get("range") or {}).get("min")
    inst.range_max = (data.get("range") or {}).get("max")
    inst.website = data.get("website")
    inst.aoop_url = data.get("aoop_url")
    inst.address = data.get("address")
    inst.save()

    _sync_relations(inst, data)
    action_log = ActionLog.objects.create(
        administrator=admin,
        action="UPDATE",
        entity="institutions",
        record_id=inst.id,
        old_data=old_data,
        new_data=data,
    )
    favorite_user_ids = list(Favorite.objects.filter(institution=inst).values_list("user_id", flat=True))
    if favorite_user_ids:
        notification_text = _build_institution_update_message(
            old_data=action_log.old_data,
            new_data=action_log.new_data,
            institution_name=inst.name,
        )
        ProfileNotification.objects.bulk_create(
            [
                ProfileNotification(
                    user_id=user_id,
                    institution=inst,
                    action_log=action_log,
                    message=notification_text,
                )
                for user_id in favorite_user_ids
            ],
            ignore_conflicts=True,
        )
    return JsonResponse({"success": True})


@csrf_exempt
@require_POST
@transaction.atomic
def delete_institution(request: HttpRequest):
    admin = _require_admin(request)
    if not admin:
        return JsonResponse({"error": "Unauthorized"}, status=401)
    data = _parse_json(request) or {}
    inst_id = data.get("id")
    if not inst_id:
        return JsonResponse({"error": "Invalid id"}, status=400)

    inst = get_object_or_404(
        Institution.objects.select_related("director").prefetch_related("conditions", "admission", "aoop_programs"),
        id=inst_id,
    )
    old_data = _serialize_institution(inst)
    director = inst.director
    inst.delete()
    if director:
        director.delete()

    ActionLog.objects.create(administrator=admin, action="DELETE", entity="institutions", record_id=inst_id, old_data=old_data)
    return JsonResponse({"success": True})
